from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE

ROOT=Path(__file__).resolve().parents[1]
DIR=ROOT/'resume'
NAVY=RGBColor(25,44,70)
TEAL=RGBColor(11,114,133)
GRAY=RGBColor(67,79,94)

def run(text,bold=False,color=None,size=None):
    pass

def para(doc,text='',style=None,space_after=0,keep=False):
    p=doc.add_paragraph(style=style)
    p.paragraph_format.space_after=Pt(space_after)
    p.paragraph_format.keep_with_next=keep
    if text:r=p.add_run(text);r.font.color.rgb=GRAY
    return p

def section(doc,title):
    p=doc.add_paragraph()
    p.paragraph_format.space_before=Pt(9);p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
    r=p.add_run(title.upper());r.bold=True;r.font.size=Pt(9.5);r.font.color.rgb=TEAL
    return p

def bullet(doc,txt):
    p=doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent=Cm(.45)
    p.paragraph_format.first_line_indent=Cm(-.2)
    p.paragraph_format.space_after=Pt(2)
    r=p.add_run(txt);r.font.size=Pt(9.1);r.font.color.rgb=GRAY
    return p

def heading2(doc,title,dates=''):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2)
    p.paragraph_format.keep_with_next=True
    r=p.add_run(title);r.bold=True;r.font.size=Pt(10.3);r.font.color.rgb=NAVY
    if dates:
        r=p.add_run('   '+dates);r.font.size=Pt(9);r.font.color.rgb=GRAY

def build(filename,role,profile,project_bullets,extra):
    d=Document()
    s=d.sections[0];s.top_margin=Cm(1.3);s.bottom_margin=Cm(1.2);s.left_margin=Cm(1.7);s.right_margin=Cm(1.7)
    s.header_distance=Cm(.5);s.footer_distance=Cm(.45)
    normal=d.styles['Normal'];normal.font.name='Aptos';normal.font.size=Pt(9.2);normal.font.color.rgb=GRAY
    normal.paragraph_format.line_spacing=1.07;normal.paragraph_format.space_after=Pt(0)
    if 'List Bullet' in d.styles:
        d.styles['List Bullet'].font.name='Aptos'
    title=d.styles['Title']
    for border in title.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.space_after=Pt(2)
    p.style=d.styles['Title']
    r=p.add_run('МАКСИМ МАКСИМОВ');r.bold=True;r.font.size=Pt(18);r.font.color.rgb=NAVY
    p=d.add_paragraph();p.paragraph_format.space_after=Pt(4)
    r=p.add_run(role);r.bold=True;r.font.size=Pt(11);r.font.color.rgb=TEAL
    p=para(d,'GitHub: https://github.com/ONN9IX',space_after=3)
    p.runs[0].font.size=Pt(8.6)

    section(d,'Профиль')
    para(d,profile,space_after=1)

    section(d,'Инструменты и методы')
    para(d,'Excel: формулы, работа с таблицами, сводные показатели;  SQL: SELECT, JOIN, GROUP BY, CTE, оконные функции;  Python: базовый уровень, Pandas;  GitHub Issues / PR;  Yandex DataLens.',space_after=2)
    para(d,'Проектная аналитика: требования, декомпозиция, приоритеты, критерии приёмки, статусная отчётность, реестр рисков, WBS.',space_after=1)

    section(d,'Практический проектный опыт')
    heading2(d,'ПРОМАКС / Smart Garden — собственный IT-проект','с 09.2026')
    for t in project_bullets:bullet(d,t)
    para(d,'Проект: github.com/ONN9IX/smart-garden',space_after=1)

    section(d,'Аналитический кейс')
    heading2(d,'Project Operations Reporting — на основе GitHub Smart Garden','10.2026')
    bullet(d,'Обработаны 137 Issues и 95 Pull Requests из публичного исторического среза; отдельно рассчитаны статусы задач и PR.')
    bullet(d,'Подготовлены Excel-отчёт с формулами, SQL-запросы, статусный отчёт, реестр рисков и модель зависимостей. Отмечены ограничения метрик и отсутствие данных для план-факт бюджета.')
    para(d,'Портфолио: https://github.com/ONN9IX/project-analyst-portfolio',space_after=1)

    section(d,'Образование')
    heading2(d,'Финансовый университет при Правительстве РФ','3 курс')
    para(d,'Направление «Бизнес-информатика».')
    if extra:
        section(d,'Дополнительно')
        para(d,extra)
    out=DIR/filename
    for p in d.paragraphs:
        text=p.text
        if 'https://github.com/' in text and len(p.runs)==1:
            before,url=text.split('https://',1)
            p.runs[0].text=before
            link=OxmlElement('w:hyperlink')
            link.set(qn('r:id'),d.part.relate_to('https://'+url,RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='https://'+url;r.append(t);link.append(r);p._p.append(link)
    d.save(out);print(out)

common=[
    'Работаю с требованиями к SaaS-платформе: выделяю функциональные сценарии, уточняю ожидаемое поведение и критерии приёмки.',
    'Структурирую задачи и изменения через GitHub Issues и Pull Requests; участвую в проверке результатов и обсуждении блокирующих факторов.',
    'Учитываю зависимости между модулями и различаю реализованные функции, отложенный scope и условия готовности пилота.'
]
build('Junior_Project_Analyst.docx','Junior Project Analyst · Аналитик IT-проектов',
      'Студент 3 курса «Бизнес-информатика». Развиваюсь в аналитике и координации IT-проектов. Работаю с Excel и SQL, знаю основы Python. Есть практика описания требований, структурирования задач, проектной отчётности и анализа данных на собственном IT-продукте.',common,
      'Рассматриваю Junior-позиции и стажировки в аналитике IT-проектов, PMO и смежном бизнес-анализе.')
build('Junior_PMO_Analyst.docx','Junior PMO Analyst · Аналитик проектного офиса',
      'Студент 3 курса «Бизнес-информатика». Интересуюсь проектным контролем, структурой задач и управленческой отчётностью. Владею Excel и SQL, Python — базовый уровень. На собственном IT-проекте работаю с этапами, критериями готовности и историей изменений.',[
      'Участвую в планировании этапов SaaS-продукта: выделяю работы, согласую зависимости и критерии готовности.',
      'Отслеживаю задачи, Pull Requests, статусы и блокеры; отделяю завершённые поставки от открытых направлений.',
      'Систематизирую проектные материалы и готовлю аналитические выводы с указанием ограничений данных.'
      ],'Интересуют позиции с Excel-отчётностью, контролем исполнения, управлением портфелем IT-задач и работой с проектными рисками.')
